import importlib.util
import json
from pathlib import Path
import re
import unittest
import zipfile

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('paper_builder',ROOT/'scripts/build_round5_paper.py')
builder=importlib.util.module_from_spec(spec);spec.loader.exec_module(builder)


class TestRound5Paper(unittest.TestCase):
    def test_revised_paper_keeps_material_limits(self):
        text='\n'.join(p.text for p in Document(builder.OUTPUT).paragraphs)
        for required in ['E3 tiếp tục là cấu hình đại diện','ngưỡng ứng viên rời rạc',
                         'chưa có MAE E1/E2','một hạt giống ngẫu nhiên (seed 42)',
                         'UCDCN chưa được chạy làm đối chứng trực tiếp']:
            self.assertIn(required,text)

    def test_tables_match_verified_metrics_and_translations(self):
        tables=builder.data_tables()
        audit=json.loads((ROOT/'reports/ROUND_5_TEST_RESULTS_2026-10-16.json').read_text(encoding='utf-8'))
        self.assertEqual(tables['results'][2][1][3],'2.7778')
        for row in tables['results'][2]:
            self.assertAlmostEqual(float(row[3])/100,audit['runs'][row[0]]['metrics']['acer'],places=6)
        for name in ['split','schedule','results','depth']:
            self.assertEqual(tables[name][2],tables[name+'_en'][2])

    def test_abstract_lengths_and_no_unfilled_text(self):
        s=(builder.PAPER/'PAPER.md').read_text(encoding='utf-8')
        for tag,nexttag,lower,upper in [('abstract_vi','abstract_en',150,200),('abstract_en','pipeline',100,150)]:
            part=s.split(f'<!-- {tag} -->')[1].split(f'<!-- {nexttag} -->')[0]
            paragraphs=[x.strip() for x in part.split('\n\n') if x.strip() and not x.strip().startswith(('##','Từ khóa','Keywords'))]
            self.assertTrue(lower<=len(paragraphs[0].split())<=upper,(tag,len(paragraphs[0].split())))
        self.assertNotRegex(s,r'\[Điền|\[CHÈN|TODO|TBD')
        self.assertIn('thay vì nội suy liên tục',s)
        self.assertIn('chưa có giá trị trung bình, độ lệch chuẩn',s.lower())

    def test_docx_contains_native_tables_and_no_template_samples(self):
        doc=Document(builder.OUTPUT)
        self.assertEqual(len(doc.tables),8)
        self.assertEqual(len(doc.sections),7)
        for s in doc.sections:
            self.assertAlmostEqual(s.page_width.cm,19,places=2)
            self.assertAlmostEqual(s.page_height.cm,26.5,places=2)
        text='\n'.join(p.text for p in doc.paragraphs)
        self.assertNotRegex(text,r'TIÊU ĐỀ TIẾNG VIỆT|Nguyễn Văn A|phím tắt|<!--|\[Điền')
        self.assertIn('E3 tiếp tục là cấu hình đại diện',text)
        self.assertEqual(doc.paragraphs[0].style.name,'Title')
        for paragraph in doc.paragraphs:
            if paragraph.text.startswith(('Bảng ', 'Table ')):
                self.assertTrue(paragraph.paragraph_format.keep_with_next)
                self.assertTrue(paragraph.paragraph_format.keep_together)
        for t,key in zip(doc.tables,['split','schedule','results','depth','split_en','schedule_en','results_en','depth_en']):
            self.assertEqual([[c.text for c in row.cells] for row in t.rows][1:],builder.data_tables()[key][2])
            for row in t.rows:
                for cell in row.cells:
                    self.assertEqual(cell.paragraphs[0].paragraph_format.first_line_indent,0)

    def test_original_template_and_preserved_parts_unchanged(self):
        provenance=json.loads((builder.PAPER/'paper_provenance.json').read_text(encoding='utf-8'))
        self.assertEqual(builder.sha(builder.TEMPLATE),provenance['template_sha256'])
        with zipfile.ZipFile(builder.TEMPLATE) as a,zipfile.ZipFile(builder.OUTPUT) as b:
            for name in a.namelist():
                if name.startswith(('word/fonts/','customXml/')) or name in ['word/numbering.xml','word/theme/theme1.xml']:
                    original, revised = a.read(name), b.read(name)
                    if name.endswith('.xml'):
                        # python-docx may change XML declaration quotes without
                        # changing the preserved template part's content.
                        original = etree.tostring(etree.fromstring(original), method='c14n')
                        revised = etree.tostring(etree.fromstring(revised), method='c14n')
                    self.assertEqual(original,revised,name)


if __name__=='__main__':unittest.main()
