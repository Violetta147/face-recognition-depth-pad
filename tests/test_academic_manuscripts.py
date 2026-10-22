"""Check the submission prose independently of internal audit documentation."""
import hashlib
import json
from pathlib import Path
import re
import unittest
import zipfile

from docx import Document
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
DOCUMENTS = {
    'paper': ROOT / 'reports/paper_round5/Nhom1_Paper_CDCN_E0_E4_20261016.docx',
    'report': ROOT / 'reports/bao_cao_cuoi_ky/ban_thao/Nhom1_Bao_cao_cuoi_ky_DeepFace_PAD_20261016.docx',
}
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}


class AcademicManuscriptTest(unittest.TestCase):
    def test_no_internal_logs_or_ambiguous_readout_language(self):
        forbidden = (
            r'SHA.?256|/content/|\.json\b|\.csv\b|\.ckpt\b|\bcommit\b|'
            r'kiểm toán|mã băm|Markdown|unittest|preflight|smoke|pilot|'
            r'artifact|persistent failure|đường dẫn|run ID|CASIA_E\d|'
            r'bộ đọc|cách đọc|đầu đọc|học từ bản đồ|nhãn suy ra|target suy ra|'
            r'ghép video|lịch học|tách người|bỏ freeze|lượt [1-8]'
        )
        for name, path in DOCUMENTS.items():
            with self.subTest(document=name), zipfile.ZipFile(path) as z:
                xml = etree.fromstring(z.read('word/document.xml'))
                text = ' '.join(xml.xpath('//w:t/text()', namespaces=NS))
                self.assertNotRegex(text.lower(), forbidden.lower())

    def test_experiment_numbers_and_equations_preserved(self):
        baseline = json.loads((ROOT / 'reports/ACADEMIC_EDIT_INVARIANTS_2026-10-19.json').read_text(encoding='utf-8'))
        for name, path in DOCUMENTS.items():
            with self.subTest(document=name):
                doc = Document(path)
                invariant = baseline['documents'][name]
                for table in invariant['numeric_tables']:
                    # The restored front-matter allocation shifts scientific
                    # tables by one without changing the historical baseline.
                    index=table['current_table_index']
                    actual = doc.tables[index + (1 if name=='report' and index>=1 else 0)]
                    values = re.findall(r'\d+(?:[.,]\d+)*', ' '.join(c.text for row in actual.rows for c in row.cells))
                    self.assertEqual(values, table['values'])
                paragraphs = [p.text for p in doc.paragraphs]
                for equation in invariant['equations']:
                    self.assertIn(equation, paragraphs)

    def test_original_styles_assets_and_hidden_breaks_preserved(self):
        baseline = json.loads((ROOT / 'reports/ACADEMIC_EDIT_INVARIANTS_2026-10-19.json').read_text(encoding='utf-8'))
        for name, path in DOCUMENTS.items():
            with self.subTest(document=name), zipfile.ZipFile(path) as z:
                invariant = baseline['documents'][name]
                for part, expected in invariant['protected_word_parts'].items():
                    self.assertEqual(hashlib.sha256(z.read(part)).hexdigest(), expected, part)
                xml = etree.fromstring(z.read('word/document.xml'))
                breaks=xml.xpath('//w:br/@w:type', namespaces=NS)
                expected=invariant['break_types']
                if name=='report':expected=[expected[0], 'page', *expected[1:]]
                self.assertEqual(breaks, expected)


if __name__ == '__main__':
    unittest.main()
