import importlib.util
import json
from pathlib import Path
import re
import statistics
import tempfile
import unittest

from docx import Document
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('round8docs',ROOT/'scripts/build_round8_documents.py')
b=importlib.util.module_from_spec(spec);spec.loader.exec_module(b)


class Round8DocumentsTest(unittest.TestCase):
    def test_new_sources_have_three_report_chapters_and_supported_scope(self):
        s=(b.WORK/'BAO_CAO.md').read_text(encoding='utf8')
        self.assertEqual(re.findall(r'^# Chương (\d)',s,re.M),['1','2','3'])
        for mode in ['report','paper']:
            text=(b.WORK/b.SOURCES[mode]).read_text(encoding='utf8')
            for phrase in ['seed 44','2,6857','0,7961','0,000684','2/5','1.080']:
                self.assertIn(phrase,text)
            self.assertNotRegex(text,r'\[Điền|\[CHÈN|TODO|TBD|artifacts/|SHA-256|\.json|\.csv|python -m')
            for rel in re.findall(r'!\[[^\]]*\]\(([^)]+)\)',text):self.assertTrue((b.WORK/rel).is_file())
        paper=(b.WORK/'PAPER.md').read_text(encoding='utf8')
        self.assertNotIn('::abstract:: Abstract',paper)
        self.assertNotIn('::keywords:: Keywords:',paper)
        for term in ['Face anti-spoofing','Presentation attack detection','Central difference convolution','Pseudo-depth','Auxiliary supervision']:
            self.assertIn(term,paper)

    def test_multiseed_rows_match_counts_and_saved_results(self):
        audit=json.loads((ROOT/'reports/ROUND_6_RESULTS_2026-10-22.json').read_text(encoding='utf8'))
        for source in b.SOURCES.values():
            rows=next(t for t in b.read_tables(b.WORK/source) if t[0][:3]==['Mô hình','Seed','APCER %'])
            self.assertEqual(len(rows),7)
            for row in rows[1:]:
                r=audit['runs'][row[0]+'_'+row[1]]
                for value,key in zip(row[2:5],['apcer','bpcer','acer']):self.assertAlmostEqual(float(value.replace(',','.')),100*r['test'][key],places=4)
                self.assertEqual(row[5:], [str(r['false_accepts'])+'/270',str(r['false_rejects'])+'/90'])
                self.assertAlmostEqual(r['test']['acer'],(r['false_accepts']/270+r['false_rejects']/90)/2)
        for model in ['E1','E3']:
            values=[audit['runs'][f'{model}_{s}']['test']['acer'] for s in [43,44,45]]
            self.assertAlmostEqual(statistics.stdev(values),audit['summary'][model]['acer']['sample_sd_ddof1'])
        delta=[audit['runs'][f'E3_{s}']['test']['acer']-audit['runs'][f'E1_{s}']['test']['acer'] for s in [43,44,45]]
        self.assertAlmostEqual(statistics.stdev(delta),audit['paired_delta']['delta_sample_sd_ddof1'])
        text=(b.WORK/'PAPER.md').read_text(encoding='utf8')
        for run in audit['runs'].values():
            self.assertIn(f"{run['validation']['threshold']:.6f}".replace('.',','),text)
        for source in b.SOURCES.values():
            rows=next(t for t in b.read_tables(b.WORK/source) if t[0]==['Chỉ số','E1','E3'])
            for row in rows[1:]:
                key=row[0].split()[0].lower()
                for model,value in zip(['E1','E3'],row[1:]):
                    factor=1 if key=='auc' else 100
                    expected=audit['summary'][model][key]
                    for number,k in zip(value.split(' ± '),['mean','sample_sd_ddof1']):
                        self.assertAlmostEqual(float(number.replace(',','.')),factor*expected[k],places=6 if key=='auc' else 4)

    def test_new_cost_and_webcam_tables_match_saved_evidence(self):
        cost=json.loads((ROOT/'reports/ROUND_6_COST_RESULTS_2026-10-22.json').read_text(encoding='utf8'))
        web=json.loads((ROOT/'reports/round7/WEBCAM_RESULTS.json').read_text(encoding='utf8'))
        for source in b.SOURCES.values():
            tables=b.read_tables(b.WORK/source)
            rows=next(t for t in tables if t[0][:3]==['Mô hình','Seed','Tham số'])
            for row in rows[1:]:
                self.assertEqual(int(row[2]),cost['models'][row[0]]['params'])
                expected=[cost['summary']['pooled'][row[0]][scope][q] for scope in ['cuda_event_forward_ms','synchronized_host_forward_ms'] for q in ['p50','p95']]
                for value,x in zip(row[3:],expected):self.assertAlmostEqual(float(value.replace(',','.')),x,places=4)
            rows=next(t for t in tables if t[0][0]=='Lượt')
            self.assertEqual(len(rows),6)
            for row,r in zip(rows[1:],web['runs']):
                for value,key in zip([row[1],row[3],row[4]],['elapsed_seconds','loop_fps','model_processed_fps']):self.assertAlmostEqual(float(value.replace(',','.')),r[key],places=3)
                if not r['first_decisions']:self.assertEqual(row[-1],'Thiếu')

    def test_builds_are_editable_and_protect_historical_outputs(self):
        old={p:b.sha(p) for p in b.HISTORICAL}
        with tempfile.TemporaryDirectory() as temp:
            for mode in ['report','paper']:
                target=b.build(mode,Path(temp)/(mode+'.docx'))
                d=Document(target)
                self.assertEqual([[[c.text for c in r.cells] for r in t.rows] for t in d.tables],b.read_tables(b.WORK/b.SOURCES[mode]))
                self.assertEqual(len(d.inline_shapes),6 if mode=='report' else 2)
                self.assertEqual(d.paragraphs[0].style.name,'Title')
                for table in d.tables:
                    self.assertTrue(all(p.paragraph_format.keep_with_next for c in table.rows[0].cells for p in c.paragraphs))
                self.assertTrue(all(p.alignment == b.WD_ALIGN_PARAGRAPH.LEFT for p in d.paragraphs if p.style.name == 'Reference'))
                self.assertNotIn('::', '\n'.join(p.text for p in d.paragraphs))
                if mode=='paper':
                    text='\n'.join(p.text for p in d.paragraphs)
                    self.assertIn('Trịnh Minh Đạt và Nguyễn Hữu Minh Quân',text)
                    self.assertIn('Khoa Công nghệ thông tin',text)
                for section in d.sections:
                    self.assertAlmostEqual(section.page_width.cm,21,places=2)
                    self.assertAlmostEqual(section.page_height.cm,29.7,places=2)
                if mode=='report':
                    self.assertEqual(len(d._element.xpath('.//m:oMath')),5)
                    self.assertGreater(len(d._element.xpath('.//m:sSub')),10)
                    self.assertEqual(len(d._element.xpath('.//m:sSup')),1)
                    self.assertNotRegex(''.join(d._element.xpath('.//m:t/text()')),r'[_^]')
                    self.assertIn('Chuong3',d._element.xml);self.assertIn('TOC',d._element.xml)
                else:
                    self.assertTrue(any(s._sectPr.find(qn('w:cols')).get(qn('w:num'))=='2' for s in d.sections))
                    self.assertEqual(d.styles['Normal'].font.size.pt,10)
                    self.assertTrue(all(s.start_type == b.WD_SECTION_START.CONTINUOUS for s in d.sections[1:]))
                with self.assertRaises(ValueError):b.build(mode,b.HISTORICAL[1])
        self.assertEqual(old,{p:b.sha(p) for p in b.HISTORICAL})


if __name__=='__main__':unittest.main()
