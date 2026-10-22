import importlib.util
import json
from pathlib import Path
import re
import tempfile
import unittest

from docx import Document
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('completed_report',ROOT/'scripts/build_final_report_completed.py')
report=importlib.util.module_from_spec(spec);spec.loader.exec_module(report)


class CompletedReportTest(unittest.TestCase):
    def test_accepted_word_preserves_tables_and_supported_claims(self):
        doc=Document(report.OUTPUT)
        self.assertEqual([[[c.text for c in r.cells] for r in t.rows] for t in doc.tables], report.read_tables())
        text='\n'.join(p.text for p in doc.paragraphs)
        for required in ['E0 có AUC cao hơn E4, đạt 0,997984', 'chưa đủ bằng chứng khẳng định E2 nhanh hơn',
                         'tổng cộng 138 chu kỳ','casia_s39_v1','không nội suy liên tục',
                         'E3 vẫn là cấu hình đại diện được lựa chọn trên tập kiểm định']:
            self.assertIn(required,text)
        self.assertNotIn('do mô hình dễ bị quá khớp',text)
        self.assertIn('hàm mất mát entropy chéo nhị phân (BCE) trên từng khung hình',text)
        self.assertNotIn('BCE) trên từng video',text)
        self.assertEqual(len(doc.inline_shapes),4)

    def test_test_results_match_frozen_audit(self):
        audit=json.loads((ROOT/'reports/ROUND_5_TEST_RESULTS_2026-10-16.json').read_text(encoding='utf-8'))
        rows=next(t for t in report.read_tables() if t[0][1]=='APCER %')
        for row in rows[1:]:
            m=audit['runs'][row[0]]['metrics']
            for value,key in zip(row[1:],['apcer','bpcer','acer','eer','auc']):
                expected=m[key]*(1 if key=='auc' else 100)
                self.assertAlmostEqual(float(value.replace(',','.')),expected,places=4 if key!='auc' else 6)

    def test_depth_and_thresholds_match_saved_validation(self):
        audit=json.loads((ROOT/'reports/ROUND_5_FULL_VALIDATION_2026-10-16.json').read_text(encoding='utf-8'))
        rows=next(t for t in report.read_tables() if t[0][1]=='Attack full')
        for row in rows[1:]:
            c=audit['runs'][row[0]]['depth_validation']['by_class']
            expected=[c['attack']['frame_mean_mae_full'],c['live']['frame_mean_mae_full'],c['live']['pixel_weighted_mae_target_mask']]
            for value,v in zip(row[1:],expected):self.assertAlmostEqual(float(value.replace(',','.')),v,places=5)
        val=next(t for t in report.read_tables() if t[0][1]=='Ngưỡng')
        for row in val[1:]:
            if row[0] in audit['runs']:
                self.assertAlmostEqual(float(row[1].replace(',','.')),audit['runs'][row[0]]['validation']['threshold'],places=6)

    def test_pair_counts_and_errors_are_consistent(self):
        rows=next(t for t in report.read_tables() if t[0][1]=='Sửa lỗi')
        for row in rows[1:]:self.assertEqual(sum(map(int,row[1:])),360)
        self.assertEqual(rows[1][1:],['0','0','5','355'])
        self.assertEqual(rows[-1][1:],['1','0','8','351'])

    def test_complete_content_and_front_member_allocation(self):
        s=report.SOURCE.read_text(encoding='utf-8')
        self.assertEqual(len(report.pages()),20)
        self.assertNotRegex(s,r'\[Điền|\[CHÈN|TODO|TBD')
        for text in ['TS. Huỳnh Hữu Hưng','Khoa Công nghệ thông tin','1021673.2610.23.14','không nội suy liên tục','tổng cộng 138 chu kỳ']:
            self.assertIn(text,s)
        self.assertLess(s.index('# Phân công trách nhiệm'),s.index('# Tóm tắt đề tài'))
        doc=Document(report.OUTPUT)
        allocation_table=doc.tables[1]
        self.assertEqual(allocation_table.cell(0,0).text,'Thành viên')
        self.assertEqual(len(allocation_table.rows),3)
        allocation_text=' '.join(c.text for row in allocation_table.rows for c in row.cells)
        for name in ['Trịnh Minh Đạt (102230180)','Nguyễn Hữu Minh Quân (102230208)']:
            self.assertIn(name,allocation_text)
        self.assertNotRegex(allocation_text,r'dẫn lượt|chưa phải|đóng góp lịch sử')
        from zipfile import ZipFile
        from lxml import etree
        with ZipFile(ROOT/'reports/paper_round5/Nhom1_Paper_CDCN_E0_E4_20261016.docx') as z:
            text=' '.join(etree.fromstring(z.read('word/document.xml')).xpath('//w:t/text()',namespaces={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}))
        self.assertNotIn('Phân công trách nhiệm',text)
        self.assertNotIn('Trách nhiệm chuyên môn',text)
        cited=set(re.findall(r'\[(\d+)\]',s.split('<!-- PAGE 19')[0]))
        references=set(re.findall(r'^\[(\d+)\]',s,re.M))
        self.assertEqual(cited,references)
        allocation=(report.WORK/'PHAN_CONG_HAI_THANH_VIEN.md').read_text(encoding='utf-8')
        self.assertEqual(len(re.findall(r'^\| [1-8] \|',allocation,re.M)),8)
        self.assertIn('50%–50%',allocation)

    def test_editable_document_and_preserved_reference(self):
        before=report.sha(report.REFERENCE)
        with tempfile.TemporaryDirectory() as temp:
            out=report.build(target=Path(temp)/'test.docx')
            d=Document(out)
            self.assertEqual(len(d.tables),14)
            self.assertEqual(len(d.inline_shapes),4)
            self.assertEqual(d.paragraphs[0].style.name,'Title')
            self.assertAlmostEqual(d.sections[0].left_margin.cm,3,places=2)
            self.assertAlmostEqual(d.sections[0].page_width.cm,21,places=2)
            xml=d._element.xml
            self.assertIn('TOC',xml);self.assertIn('PAGEREF T12',xml)
            self.assertEqual(len(d._element.xpath('.//w:bookmarkStart')),16)
            # Front matter and seven major parts start separately; internal
            # content chunks flow, so prose edits do not leave stale page gaps.
            self.assertEqual(len(d._element.xpath('.//w:br[@w:type="page"]')),10)
            for table in d.tables:
                for row in table.rows:
                    self.assertIsNotNone(row._tr.trPr.find(qn('w:cantSplit')))
            self.assertEqual(report.sha(report.REFERENCE),before)


if __name__=='__main__':unittest.main()
