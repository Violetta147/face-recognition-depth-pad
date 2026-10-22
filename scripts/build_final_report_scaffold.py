import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def field(paragraph, instruction):
    run = paragraph.add_run()
    start = OxmlElement("w:fldChar")
    start.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (start, text, end):
        run._r.append(element)


def set_font(style, size, bold=False):
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attribute in list(fonts.attrib):
        if "theme" in attribute.lower():
            del fonts.attrib[attribute]
    for name in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{name}"), "Times New Roman")
    size_cs = style.element.get_or_add_rPr().find(qn("w:szCs"))
    if size_cs is not None:
        size_cs.set(qn("w:val"), str(size * 2))
    properties = style.element.find(qn("w:pPr"))
    if properties is not None:
        for border in list(properties.findall(qn("w:pBdr"))):
            properties.remove(border)


def add_table(document, lines):
    rows = [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in lines]
    rows = [row for row in rows if not all(re.fullmatch(r":?-+:?", cell) for cell in row)]
    count = len(rows[0])
    if any(len(row) != count for row in rows):
        raise ValueError("Table rows have different column counts")
    table = document.add_table(rows=0, cols=count)
    table.autofit = False
    fractions = {2: [0.5, 0.5], 3: [0.27, 0.43, 0.3], 4: [0.28, 0.24, 0.24, 0.24],
                 5: [0.24, 0.16, 0.16, 0.16, 0.28], 6: [0.35, 0.13, 0.13, 0.13, 0.13, 0.13]}[count]
    for column, fraction in zip(table.columns, fractions):
        column.width = Cm(16 * fraction)
    for index, values in enumerate(rows):
        cells = table.add_row().cells
        for cell, content, fraction in zip(cells, values, fractions):
            cell.width = Cm(16 * fraction)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            properties = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side in ("top", "bottom", "left", "right"):
                element = OxmlElement(f"w:{side}")
                element.set(qn("w:w"), "85")
                element.set(qn("w:type"), "dxa")
                margins.append(element)
            properties.append(margins)
            borders = OxmlElement("w:tcBorders")
            for side in ("top", "bottom", "left", "right"):
                element = OxmlElement(f"w:{side}")
                for key, value in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
                    element.set(qn(f"w:{key}"), value)
                borders.append(element)
            properties.append(borders)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "E7E6E6" if index == 0 else "FFFFFF")
            properties.append(shading)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.1
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(content)
            run.font.size = Pt(11)
            run.bold = index == 0
        if index == 0:
            table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def build(source, target):
    markdown = source.read_text(encoding="utf-8")
    chunks = re.split(r"<!-- PAGE (\d+) \| ([^>]+) -->", markdown)
    numbers = [int(chunks[i]) for i in range(1, len(chunks), 3)]
    if numbers != list(range(1, 19)):
        raise ValueError("Expected 18 ordered scaffold pages")
    document = Document()
    document.core_properties.title = "Khung báo cáo cuối kỳ Thị giác máy tính"
    document.core_properties.author = "Trịnh Minh Đạt; Nguyễn Hữu Minh Quân"
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin, section.right_margin = Cm(3), Cm(2)
    section.top_margin = section.bottom_margin = Cm(2)
    section.footer_distance = Cm(1)
    section.different_first_page_header_footer = True
    styles = document.styles
    set_font(styles["Normal"], 13)
    styles["Normal"].paragraph_format.line_spacing = 1.3
    styles["Normal"].paragraph_format.space_after = Pt(6)
    set_font(styles["Title"], 18, True)
    set_font(styles["Heading 1"], 16, True)
    set_font(styles["Heading 2"], 14, True)
    for name in ("Title", "Heading 1", "Heading 2"):
        styles[name].paragraph_format.space_before = Pt(10)
        styles[name].paragraph_format.space_after = Pt(8)
        styles[name].paragraph_format.keep_with_next = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Trang ").font.size = Pt(10)
    field(footer, "PAGE")
    for offset in range(1, len(chunks), 3):
        page = int(chunks[offset])
        if page > 1:
            document.add_page_break()
        lines = chunks[offset + 2].strip().splitlines()
        index = 0
        while index < len(lines):
            line = lines[index].strip()
            index += 1
            if not line or line.startswith("<!--"):
                continue
            if line.startswith("|"):
                table_lines = [line]
                while index < len(lines) and lines[index].strip().startswith("|"):
                    table_lines.append(lines[index].strip())
                    index += 1
                add_table(document, table_lines)
                continue
            if line == "[[TOC]]":
                field(document.add_paragraph(), 'TOC \\o "1-1" \\h \\z \\u')
                continue
            if line.startswith("# "):
                paragraph = document.add_paragraph(line[2:], style="Title" if page <= 3 else "Heading 1")
            elif line.startswith("## "):
                paragraph = document.add_paragraph(line[3:], style="Heading 2")
            else:
                paragraph = document.add_paragraph(line)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if page == 1:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_after = Pt(14)
            if line.startswith("[CHÈN HÌNH"):
                paragraph.paragraph_format.space_before = Pt(12)
                paragraph.paragraph_format.space_after = Pt(12)
                paragraph.runs[0].italic = True
    settings = document.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)
    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)
    print(target)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("target", type=Path)
    args = parser.parse_args()
    build(args.source, args.target)
