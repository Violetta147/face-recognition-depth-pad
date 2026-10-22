"""Build the completed course report while preserving the original scaffold."""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / "reports/bao_cao_cuoi_ky"
SOURCE = WORK / "ban_thao/BAO_CAO_CUOI_KY.md"
REFERENCE = WORK / "ban_thao/BAO_CAO_CUOI_KY_KHUNG.docx"
OUTPUT = WORK / "ban_thao/Nhom1_Bao_cao_cuoi_ky_DeepFace_PAD_20261016.docx"


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def pages(source=SOURCE):
    chunks = re.split(r"<!-- PAGE (\d+) \| ([^>]+) -->", Path(source).read_text(encoding="utf-8"))
    result = [(int(chunks[i]), chunks[i+1].strip(), chunks[i+2].strip()) for i in range(1,len(chunks),3)]
    if [x[0] for x in result] != list(range(1,21)):
        raise ValueError("Expected twenty ordered report content sections")
    return result


def field(paragraph, instruction):
    run=paragraph.add_run()
    start=OxmlElement('w:fldChar');start.set(qn('w:fldCharType'),'begin')
    text=OxmlElement('w:instrText');text.set(qn('xml:space'),'preserve');text.text=instruction
    separate=OxmlElement('w:fldChar');separate.set(qn('w:fldCharType'),'separate')
    end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end')
    for item in [start,text,separate,end]:run._r.append(item)


def font(run,size=13,bold=False):
    run.font.name='Times New Roman';run.font.size=Pt(size)
    run.font.bold=bold;run.font.color.rgb=RGBColor(0,0,0)
    rf=run._r.get_or_add_rPr().get_or_add_rFonts()
    for key in list(rf.attrib):
        if 'theme' in key.lower():del rf.attrib[key]
    for key in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+key),'Times New Roman')


def bookmark(p,name,number):
    start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(number));start.set(qn('w:name'),name)
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(number))
    p._p.insert(0,start);p._p.append(end)


def captions(source=SOURCE):
    text=Path(source).read_text(encoding='utf-8')
    tables=re.findall(r'^(Bảng (\d+) .+)$',text,re.M)
    figures=re.findall(r'^!\[(Hình (\d+) [^\]]+)\]\(([^)]+)\)$',text,re.M)
    return tables,figures


def read_tables(source=SOURCE):
    return [[[x.strip() for x in line.strip('|').split('|')] for line in block.splitlines()
             if not re.fullmatch(r'\|[\s:|\-]+',line)]
            for block in re.findall(r'(?:^\|.*\n)+',Path(source).read_text(encoding='utf-8'),re.M)]


def table(doc,rows):
    n=len(rows[0])
    if any(len(r)!=n for r in rows):raise ValueError('Inconsistent table width')
    # Numeric six-column result tables keep all metrics readable; prose gets a wider field.
    if n==6 and rows[0][1]=='APCER %':weights=[.8,1.1,1.1,1.1,1.1,1.5]
    elif n==6 and rows[0][1]=='Lịch epoch':weights=[.6,2.2,.9,.6,.9,.9]
    elif n==3 and rows[0][0]=='Thành viên':weights=[1.1,2.3,1.6]
    elif n==3 and rows[0][0]=='Công trình':weights=[1,2,2]
    elif n==4 and rows[0][0]=='Hạng mục':weights=[1.3,.7,1.5,1.5]
    elif n==3 and rows[0][0]=='Vị trí':weights=[1.2,1.4,1.6]
    elif n==3 and rows[0][0]=='Video':weights=[1.6,1,1]
    else:weights=[1]*n
    widths=[16*x/sum(weights) for x in weights]
    t=doc.add_table(rows=0,cols=n);t.autofit=False
    for column,w in zip(t.columns,widths):column.width=Cm(w)
    for i,values in enumerate(rows):
        row=t.add_row()
        for cell,value,w in zip(row.cells,values,widths):
            cell.width=Cm(w);cell.vertical_alignment=WD_ALIGN_VERTICAL.CENTER
            p=cell.paragraphs[0]
            p.paragraph_format.first_line_indent=Cm(0)
            p.paragraph_format.left_indent=p.paragraph_format.right_indent=Cm(0)
            p.paragraph_format.space_before=p.paragraph_format.space_after=Pt(2)
            p.paragraph_format.line_spacing=1.1
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT if n==3 or (n==4 and rows[0][0]=='Hạng mục') else WD_ALIGN_PARAGRAPH.CENTER
            font(p.add_run(value),11,i==0)
            props=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'60');e.set(qn('w:type'),'dxa');margins.append(e)
            props.append(margins)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7E6E6' if i==0 else 'FFFFFF');props.append(shade)
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    borders=OxmlElement('w:tblBorders')
    for side in ['top','bottom','left','right','insideH','insideV']:
        e=OxmlElement('w:'+side)
        for k,v in {'val':'single','sz':'4','color':'D9D9D9'}.items():e.set(qn('w:'+k),v)
        borders.append(e)
    t._tbl.tblPr.append(borders)
    after=doc.add_paragraph();after.paragraph_format.space_after=Pt(2)
    after.paragraph_format.line_spacing=1;font(after.add_run(''),2)


def configure(doc):
    for section in doc.sections:
        section.page_width,section.page_height=Cm(21),Cm(29.7)
        section.left_margin,section.right_margin=Cm(3),Cm(2)
        section.top_margin=section.bottom_margin=Cm(2)
        section.different_first_page_header_footer=True
    for name,size in [('Normal',13),('Title',18),('Heading 1',16),('Heading 2',14),('Caption',11),('Report Small',10),('Report Index',11),('Report Reference',11),('Report Equation',12)]:
        if name not in doc.styles:doc.styles.add_style(name,1)
        style=doc.styles[name]
        style.font.name='Times New Roman';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
        style.font.bold=name in ['Title','Heading 1','Heading 2']
        style.paragraph_format.line_spacing=1.3 if name=='Normal' else 1.1
        style.paragraph_format.space_after=Pt(6 if name=='Normal' else 4)
        style.paragraph_format.space_before=Pt(8 if name in ['Heading 1','Heading 2'] else 0)
        style.paragraph_format.widow_control=True
        style.paragraph_format.keep_with_next=name in ['Title','Heading 1','Heading 2']
        rf=style.element.get_or_add_rPr().get_or_add_rFonts()
        for key in list(rf.attrib):
            if 'theme' in key.lower():del rf.attrib[key]
        for key in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+key),'Times New Roman')
        for b in list(style.element.xpath('.//w:pBdr')):b.getparent().remove(b)
    for name in ['TOC 1','TOC 2','toc 1','toc 2']:
        if name in doc.styles:
            s=doc.styles[name];s.font.name='Times New Roman';s.font.size=Pt(11);s.font.color.rgb=RGBColor(0,0,0)
            s.paragraph_format.space_after=Pt(2);s.paragraph_format.line_spacing=1.0


def build(source=SOURCE,target=OUTPUT,reference=REFERENCE):
    source,target,reference=map(Path,[source,target,reference])
    if target.resolve()==reference.resolve():raise ValueError('Preserve the original scaffold')
    original_hash=sha(reference)
    doc=Document(reference)
    for node in list(doc._element.body):
        if node.tag!=qn('w:sectPr'):doc._element.body.remove(node)
    configure(doc)
    for sec in doc.sections:
        for h in [sec.header,sec.first_page_header,sec.footer,sec.first_page_footer]:
            for child in list(h._element):h._element.remove(child)
            h.add_paragraph()
        footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        font(footer.add_run('Trang '),10);field(footer,'PAGE')
    all_tables,all_figures=captions(source)
    image_inputs=[]
    for number,_,content in pages(source):
        # PAGE labels identify historical content chunks. Longer revised prose
        # must flow within chapters rather than inherit a break every chunk.
        starts_major_part = content.startswith(('# Chương ', '# Kết luận', '# Tài liệu tham khảo', '# Phụ lục'))
        if number in (2, 3) or starts_major_part:doc.add_page_break()
        lines=content.splitlines();i=0
        while i<len(lines):
            line=lines[i].strip();i+=1
            if not line or line.startswith('<!--'):continue
            if line=='[[PAGE_BREAK]]':
                doc.add_page_break();continue
            if line.startswith('|'):
                raw=[line]
                while i<len(lines) and lines[i].strip().startswith('|'):
                    raw.append(lines[i].strip());i+=1
                rows=[[c.strip() for c in l.strip('|').split('|')] for l in raw if not re.fullmatch(r'\|[\s:|\-]+',l)]
                table(doc,rows);continue
            if line=='[[TOC]]':
                field(doc.add_paragraph(),r'TOC \o "1-1" \h \z \u');continue
            if line in ['[[FIGURES]]','[[TABLES]]']:
                items=[(cap,n,'F') for cap,n,_ in all_figures] if line=='[[FIGURES]]' else [(cap,n,'T') for cap,n in all_tables]
                for cap,n,prefix in items:
                    p=doc.add_paragraph(style='Report Index')
                    p.paragraph_format.tab_stops.add_tab_stop(Cm(16),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
                    font(p.add_run(cap+'\t'),11);field(p,'PAGEREF '+prefix+n+r' \h')
                continue
            m=re.fullmatch(r'!\[(Hình (\d+) [^\]]+)\]\(([^)]+)\)',line)
            if m:
                cap,n,image_name=m.groups();image_path=(source.parent/image_name).resolve()
                if not image_path.is_file():raise FileNotFoundError(image_path)
                p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next=True
                p.add_run().add_picture(str(image_path),width=Cm(16))
                p=doc.add_paragraph(style='Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                font(p.add_run(cap),11);bookmark(p,'F'+n,100+int(n))
                image_inputs.append({'path':str(image_path.relative_to(ROOT)),'sha256':sha(image_path)})
                continue
            if re.match(r'^Bảng \d+ ',line):
                p=doc.add_paragraph(style='Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next=True;font(p.add_run(line),11)
                n=re.match(r'Bảng (\d+)',line).group(1);bookmark(p,'T'+n,200+int(n));continue
            if line.startswith('# '):
                p=doc.add_paragraph(style='Title' if number<=3 else 'Heading 1');text=line[2:]
            elif line.startswith('## '):
                p=doc.add_paragraph(style='Heading 2');text=line[3:]
            elif line.startswith('::small:: '):
                p=doc.add_paragraph(style='Report Small');text=line[len('::small:: '):]
            elif line.startswith('::equation:: '):
                p=doc.add_paragraph(style='Report Equation');text=line[len('::equation:: '):]
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            else:
                p=doc.add_paragraph(style='Report Reference' if number==19 else 'Normal');text=line
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if number in [19,20] else WD_ALIGN_PARAGRAPH.JUSTIFY
            font(p.add_run(text),p.style.font.size.pt,p.style.font.bold or False)
            if text=='Phân công trách nhiệm của hai thành viên':
                p.paragraph_format.page_break_before=True
            if number==1:
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after=Pt(12)
    setting=doc.settings.element.find(qn('w:updateFields'))
    if setting is None:setting=OxmlElement('w:updateFields');doc.settings.element.append(setting)
    setting.set(qn('w:val'),'true')
    doc.core_properties.title='Báo cáo cuối kỳ DeepFace PAD đến lượt 5'
    doc.core_properties.author='Trịnh Minh Đạt; Nguyễn Hữu Minh Quân'
    doc.core_properties.subject='Thị giác máy tính 1021673.2610.23.14'
    target.parent.mkdir(parents=True,exist_ok=True);doc.save(target)
    assert sha(reference)==original_hash
    output_label=str(target.relative_to(ROOT)) if target.is_relative_to(ROOT) else str(target)
    provenance={'source_sha256':sha(source),'reference_sha256':original_hash,'output_sha256_before_field_refresh':sha(target),
                'source':str(source.relative_to(ROOT)),'output':output_label,'images':image_inputs,
                'evidence_sources':['reports/RESEARCH_AUDIT_ROUNDS_1_4_2026-10-16.json','reports/ROUND_5_FULL_VALIDATION_2026-10-16.json','reports/ROUND_5_TEST_RESULTS_2026-10-16.json'],
                'scope':'single seed 42 CASIA development to Round 5; Round 6-8 are plans',
                'allocation':'member responsibilities in report front matter; detailed planning retained separately',
                'teacher_source':'https://scv.udn.vn/hhhung','teacher_verified_title':'TS. Huỳnh Hữu Hưng'}
    provenance_dir=WORK if target.resolve()==OUTPUT.resolve() else target.parent
    (provenance_dir/'REPORT_PROVENANCE.json').write_text(json.dumps(provenance,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(target)
    return target


if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--output',type=Path,default=OUTPUT)
    args=p.parse_args();build(target=args.output)
