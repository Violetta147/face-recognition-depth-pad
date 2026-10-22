"""Build an editable research manuscript from verified Round 5 artifacts."""
from __future__ import annotations

import argparse
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re
import zipfile

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
PAPER = ROOT / "reports/paper_round5"
TEMPLATE = ROOT / "artifacts/workspace-audit-20261016/Template_UD_JST_original.docx"
OUTPUT = PAPER / "Nhom1_Paper_CDCN_E0_E4_20261016.docx"


def pipeline_figure():
    """Scientific schematic of implemented modules, with English labels."""
    import math
    from PIL import Image, ImageDraw, ImageFont
    out=ROOT/'artifacts/paper-round5-20261016/pipeline-en.png'
    out.parent.mkdir(parents=True,exist_ok=True)
    im = Image.new('RGB', (2010, 690), 'white')
    draw = ImageDraw.Draw(im)
    face = ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 32)
    def point(x, y):
        return (round(x*2010), round((1-y)*690))
    def label(x, y, text):
        draw.multiline_text(point(x,y),text,font=face,fill='#222222',anchor='mm',align='center',spacing=8)
    def box(x,y,w,h,text):
        draw.rectangle([point(x,y+h),point(x+w,y)],fill='#F5F5F5',outline='#333333',width=3)
        label(x+w/2,y+h/2,text)
    def arrow(a,b,dashed=False):
        a,b=point(*a),point(*b)
        dx,dy=b[0]-a[0],b[1]-a[1]
        length=math.hypot(dx,dy)
        ux,uy=dx/length,dy/length
        if dashed:
            for start in range(0,round(length),22):
                end=min(start+12,length)
                draw.line((a[0]+ux*start,a[1]+uy*start,a[0]+ux*end,a[1]+uy*end),fill='#333333',width=3)
        else:
            draw.line((a,b),fill='#333333',width=3)
        draw.polygon([b,(b[0]-16*ux+7*uy,b[1]-16*uy-7*ux),(b[0]-16*ux-7*uy,b[1]-16*uy+7*ux)],fill='#333333')
    box(.01,.59,.15,.23,'RGB\n256 × 256')
    box(.22,.59,.16,.23,'CDCN\nθ = 0.7')
    box(.44,.59,.16,.23,'Depth map\n1 × 32 × 32')
    box(.68,.77,.31,.18,'E1 spatial mean')
    box(.68,.45,.31,.23,'E2 / E3 / E3b / E4\nLearned head')
    for a,b in [((.16,.705),(.22,.705)),((.38,.705),(.44,.705)),((.60,.75),(.68,.86)),((.60,.64),(.68,.56))]:arrow(a,b)
    box(.01,.08,.31,.22,'Offline pseudo-depth\n3DDFA V2 live / zero attack')
    box(.40,.08,.27,.22,'Depth MSE + contrast MSE\nBCE or Focal')
    arrow((.32,.19),(.40,.19),True);arrow((.52,.59),(.52,.30),True)
    arrow((.73,.45),(.63,.30),True)
    label(.835,.17,'20 frames → video mean\nValidation threshold')
    im.save(out,dpi=(300,300))
    return out


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def data_tables(root=ROOT):
    test = json.loads((root / "reports/ROUND_5_TEST_RESULTS_2026-10-16.json").read_text(encoding="utf-8"))
    val = json.loads((root / "reports/ROUND_5_FULL_VALIDATION_2026-10-16.json").read_text(encoding="utf-8"))
    split = [["Train", "16", "192", "3840"], ["Val", "4", "48", "960"], ["Test", "30", "360", "7200"]]
    results = []
    for name, run in test["runs"].items():
        m = run["metrics"]
        results.append([name] + [f"{100*m[k]:.4f}" for k in ("apcer", "bpcer", "acer", "eer")] + [f"{m['auc']:.6f}"])
    schedule = [["E0", "5", "BCE", "5"], ["E1", "30", "Depth", "25"],
                ["E2", "10 head", "BCE", "1"], ["E3", "16 joint", "Depth+BCE", "10"],
                ["E3b", "4+4+8", "Depth+BCE", "10"], ["E4", "4+4+8", "Depth+Focal", "14"]]
    depth = []
    for name, run in val["runs"].items():
        c = run["depth_validation"]["by_class"]
        depth.append([name, f"{c['attack']['frame_mean_mae_full']:.5f}",
                      f"{c['live']['frame_mean_mae_full']:.5f}", f"{c['live']['pixel_weighted_mae_target_mask']:.5f}"])
    return {
        "split": ("Bảng 1. Thống kê dữ liệu theo từng phân tập", ["Tập", "Người", "Video", "Frame"], split),
        "schedule": ("Bảng 2. Số chu kỳ huấn luyện và chu kỳ được chọn", ["Model", "Số chu kỳ", "Loss", "Best"], schedule),
        "results": ("Bảng 3. Kết quả trên video kiểm tra (tỷ lệ lỗi: %)", ["Model", "APCER", "BPCER", "ACER", "EER", "AUC"], results),
        "depth": ("Bảng 4. MAE độ sâu trên tập kiểm định", ["Model", "Attack full", "Live full", "Live mask"], depth),
        "split_en": ("Table 1. Dataset statistics by split", ["Split", "Subjects", "Videos", "Frames"], split),
        "schedule_en": ("Table 2. Epoch budgets and selected global epochs", ["Model", "Schedule", "Loss", "Best"], schedule),
        "results_en": ("Table 3. Video-level performance on the development test set (error rates in %)", ["Model", "APCER", "BPCER", "ACER", "EER", "AUC"], results),
        "depth_en": ("Table 4. Validation depth MAE against pseudo-depth targets", ["Model", "Attack full", "Live full", "Live mask"], depth),
    }


def font(run, size=10, bold=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:" + key), "Times New Roman")


def paragraph(doc, text, style="Paper Body", align=None, size=10, bold=False):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    font(p.add_run(text), size, bold)
    return p


def set_section(section, columns):
    section.page_width, section.page_height = Cm(19), Cm(26.5)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Cm(1)
    section.header_distance = section.footer_distance = Cm(.35)
    cols = section._sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols"); section._sectPr.append(cols)
    cols.clear()
    cols.set(qn("w:num"), str(columns)); cols.set(qn("w:space"), "283")
    cols.set(qn("w:equalWidth"), "1")


def table(doc, spec, full=False):
    caption, headers, rows = spec
    cp = paragraph(doc, caption, "Caption", size=9)
    cp.paragraph_format.keep_with_next = True
    cp.paragraph_format.keep_together = True
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    total = 17 if full else 8.25
    weights = [0.8, 1.05, 1.05, 1.05, 1.05, 1.6] if len(headers) == 6 else [1]*len(headers)
    widths = [total*x/sum(weights) for x in weights]
    for c, width in zip(t.columns, widths): c.width = Cm(width)
    for target, values in [(t.rows[0], headers)] + [(t.add_row(), values) for values in rows]:
        for j, (cell, value) in enumerate(zip(target.cells, values)):
            cell.width = Cm(widths[j]); cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.left_indent = p.paragraph_format.right_indent = Cm(0)
            p.paragraph_format.keep_with_next = target._tr is not t.rows[-1]._tr
            font(p.add_run(str(value)), 9, target._tr is t.rows[0]._tr)
            pr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side in ["top", "bottom", "left", "right"]:
                e = OxmlElement("w:"+side); e.set(qn("w:w"), "35"); e.set(qn("w:type"), "dxa"); margins.append(e)
            pr.append(margins)
            if target._tr is t.rows[0]._tr:
                shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "EDEDED"); pr.append(shade)
        cant = OxmlElement("w:cantSplit"); target._tr.get_or_add_trPr().append(cant)
    header = OxmlElement("w:tblHeader"); t.rows[0]._tr.get_or_add_trPr().append(header)
    borders = OxmlElement("w:tblBorders")
    for side in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        e = OxmlElement("w:"+side); e.set(qn("w:val"), "single");e.set(qn("w:sz"), "4");e.set(qn("w:color"), "D9D9D9"); borders.append(e)
    t._tbl.tblPr.append(borders)
    paragraph(doc, "", size=2).paragraph_format.space_after = Pt(1)


def build(source=PAPER / "PAPER.md", output=OUTPUT, template=TEMPLATE):
    source, output, template = Path(source), Path(output), Path(template)
    original_hash = sha(template)
    doc = Document(template)
    # Keep the reference's styles/theme/opaque parts, replace its example body.
    for child in list(doc._element.body):
        if child.tag != qn("w:sectPr"): doc._element.body.remove(child)
    for section in doc.sections:
        for part in [section.header, section.even_page_header, section.first_page_header,
                     section.footer, section.even_page_footer, section.first_page_footer]:
            for child in list(part._element): part._element.remove(child)
            part.add_paragraph()
    for name in ["Paper Body", "Paper Heading", "Paper Subheading"]:
        if name not in doc.styles: doc.styles.add_style(name, 1)
        s = doc.styles[name]; s.base_style = doc.styles["normal"]
        s.font.name = "Times New Roman";s.font.size=Pt(10);s.font.color.rgb=RGBColor(0,0,0)
        s.paragraph_format.space_before=s.paragraph_format.space_after=Pt(3)
        s.paragraph_format.line_spacing=1
        s.paragraph_format.widow_control=True
        s.paragraph_format.first_line_indent=Cm(.4) if name == "Paper Body" else Cm(0)
        if name != "Paper Body":
            s.font.bold=True;s.paragraph_format.keep_with_next=True
            s.paragraph_format.space_before=Pt(6)
    for name in ["Title", "Caption"]:
        if name not in doc.styles:
            doc.styles.add_style(name, 1).base_style = doc.styles["normal"]
        s=doc.styles[name];s.font.name="Times New Roman";s.font.color.rgb=RGBColor(0,0,0)
        s.paragraph_format.space_before=s.paragraph_format.space_after=Pt(3)
        s.paragraph_format.line_spacing=1
        for b in list(s._element.xpath(".//w:pBdr")): b.getparent().remove(b)
    set_section(doc.sections[0],1)
    mode="title_vi"; tables=data_tables(); buffer=[]; image_path=pipeline_figure()
    def flush():
        if not buffer:return
        text=" ".join(buffer);buffer.clear()
        if text.startswith("# "):text=text[2:]
        if text.startswith("## ") or text.startswith("### "):
            level=3 if text.startswith("### ") else 2
            p=paragraph(doc,text[level+1:],"Paper Subheading" if level==3 else "Paper Heading",bold=True)
        elif mode=="title_vi":p=paragraph(doc,text,"Title",WD_ALIGN_PARAGRAPH.CENTER,13,True)
        elif mode=="title_en":p=paragraph(doc,text,"Title",WD_ALIGN_PARAGRAPH.CENTER,11,False)
        elif mode=="authors":p=paragraph(doc,text,align=WD_ALIGN_PARAGRAPH.CENTER)
        else:
            p=paragraph(doc,text,size=9 if text.startswith("[") and mode=="body" else 10)
            p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    for line in source.read_text(encoding="utf-8").splitlines():
        m=re.fullmatch(r"<!-- (.*?) -->",line.strip())
        if m:
            flush();tag=m.group(1)
            if tag.startswith("table:"):
                key=tag.split(":")[1];table(doc,tables[key])
                continue
            mode=tag
            if tag=="abstract_vi":set_section(doc.add_section(WD_SECTION_START.CONTINUOUS),2)
            elif tag=="abstract_en":doc.paragraphs[-1].add_run().add_break(WD_BREAK.COLUMN)
            elif tag=="pipeline":
                set_section(doc.add_section(WD_SECTION_START.CONTINUOUS),1)
                p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(image_path),width=Cm(16))
                paragraph(doc,"Hình 1. Quy trình đã triển khai; đường nét liền biểu thị luồng suy luận, đường nét đứt biểu thị luồng giám sát khi huấn luyện.", "Caption",WD_ALIGN_PARAGRAPH.CENTER,9)
            elif tag=="body":set_section(doc.add_section(WD_SECTION_START.CONTINUOUS),2)
            elif tag=="appendix":
                # Close and balance the body columns before starting the appendix.
                set_section(doc.add_section(WD_SECTION_START.CONTINUOUS),1)
                set_section(doc.add_section(WD_SECTION_START.NEW_PAGE),2)
            elif tag=="pipeline_en":
                set_section(doc.add_section(WD_SECTION_START.CONTINUOUS),1)
                p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(image_path),width=Cm(16))
                paragraph(doc,"Figure 1. Overview of the implemented pipeline; solid arrows denote inference, dashed arrows indicate training supervision.", "Caption",WD_ALIGN_PARAGRAPH.CENTER,9)
            continue
        if line.strip():buffer.append(line.strip())
        else:flush()
    flush()
    for section in doc.sections:
        for footer in (section.footer, section.even_page_footer):
            p=footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
            for child in list(p._p):
                if child.tag!=qn('w:pPr'):p._p.remove(child)
            field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
    doc.core_properties.title="Đánh giá thực nghiệm bộ phân loại bản đồ độ sâu và chiến lược huấn luyện CDCN trên CASIA-FASD"
    doc.core_properties.author="Trịnh Minh Đạt; Nguyễn Hữu Minh Quân"
    doc.core_properties.subject="Course review manuscript using frozen single-seed CASIA development results"
    output.parent.mkdir(parents=True,exist_ok=True)
    doc.save(output)
    # Preserve all source parts outside the declared editable OOXML slots.
    allowed={"word/document.xml","word/styles.xml","word/_rels/document.xml.rels","docProps/core.xml","[Content_Types].xml"}
    tmp=output.with_suffix('.tmp.docx')
    with zipfile.ZipFile(template) as zref,zipfile.ZipFile(output) as znew,zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as zout:
        for item in znew.infolist():
            editable=item.filename in allowed or re.fullmatch(r"word/(?:header|footer)\d+\.xml",item.filename)
            content=znew.read(item.filename)
            if not editable and item.filename in zref.namelist():content=zref.read(item.filename)
            zout.writestr(item,content)
    tmp.replace(output)
    assert sha(template)==original_hash
    provenance={"template_sha256":original_hash,"source_sha256":sha(source),"output_sha256":sha(output),
                "tables":list(tables),"sources":["ROUND_5_TEST_RESULTS_2026-10-16.json","ROUND_5_FULL_VALIDATION_2026-10-16.json"],
                "intentional_template_changes":["replace sample body","remove publication placeholder headers and author biographies", "reuse two-column pattern for English appendix and full-width pattern for its figure", "regenerate implementation schematic with English labels; source figure unchanged"]}
    (PAPER/'paper_provenance.json').write_text(json.dumps(provenance,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(output)
    return output


if __name__=="__main__":
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output',type=Path,default=OUTPUT)
    args=parser.parse_args();build(output=args.output)
