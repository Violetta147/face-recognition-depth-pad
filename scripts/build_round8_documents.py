"""Build new Round 8 manuscripts without writing any historical deliverable.

Uses existing evidence and images only. No model, webcam or training imports.
Run with the bundled Python runtime; see the new package README for rendering.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / 'reports/round8_documents_20261023'
OUTPUTS = {'report': 'Nhom1_Bao_cao_DeepFace_PAD_20261023.docx',
           'paper': 'Nhom1_Paper_CDCN_IEEE_VI_20261023.docx'}
SOURCES = {'report': 'BAO_CAO.md', 'paper': 'PAPER.md'}
HISTORICAL = [ROOT/'reports/bao_cao_cuoi_ky/ban_thao/BAO_CAO_CUOI_KY.md',
              ROOT/'reports/bao_cao_cuoi_ky/ban_thao/Nhom1_Bao_cao_cuoi_ky_DeepFace_PAD_20261016.docx',
              ROOT/'reports/paper_round5/PAPER.md',
              ROOT/'reports/paper_round5/Nhom1_Paper_CDCN_E0_E4_20261016.docx']
EVIDENCE = ['reports/RESEARCH_AUDIT_ROUNDS_1_4_2026-10-16.json',
            'reports/ROUND_5_FULL_VALIDATION_2026-10-16.json',
            'reports/ROUND_5_TEST_RESULTS_2026-10-16.json',
            'reports/ROUND_6_RESULTS_2026-10-22.json',
            'reports/ROUND_6_COST_RESULTS_2026-10-22.json',
            'reports/ROUND_6_PAPER_PROTOCOL_COMPARISON_2026-10-22.json',
            'reports/round7/analysis/analysis.json',
            'reports/round7/analysis/subgroups.csv',
            'reports/round7/figures-attempt03/provenance.json',
            'reports/round7/FIGURES_QA.json',
            'reports/round7/WEBCAM_RESULTS.json']


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_tables(source):
    return [[[c.strip() for c in line.strip('|').split('|')]
             for line in block.splitlines() if not re.fullmatch(r'\|[\s:|\-]+', line)]
            for block in re.findall(r'(?:^\|.*\n)+', Path(source).read_text(encoding='utf8'), re.M)]


def field(p, instruction):
    run = p.add_run()
    for tag, value in [('begin', None), ('instruction', instruction), ('separate', None), ('end', None)]:
        e = OxmlElement('w:instrText' if tag == 'instruction' else 'w:fldChar')
        if value is not None:
            e.set(qn('xml:space'), 'preserve'); e.text = value
        else:
            e.set(qn('w:fldCharType'), tag)
        run._r.append(e)


def bookmark(p, name, index):
    start = OxmlElement('w:bookmarkStart'); start.set(qn('w:id'), str(index)); start.set(qn('w:name'), name)
    end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), str(index))
    p._p.insert(0, start); p._p.append(end)


def font(run, size, bold=False):
    run.font.name = 'Times New Roman'; run.font.size = Pt(size)
    run.font.bold = bold; run.font.color.rgb = RGBColor(0, 0, 0)
    rf = run._r.get_or_add_rPr().get_or_add_rFonts()
    for key in list(rf.attrib):
        if 'theme' in key.lower(): del rf.attrib[key]
    for key in ['ascii', 'hAnsi', 'eastAsia', 'cs']: rf.set(qn('w:'+key), 'Times New Roman')


def columns(section, count):
    node = section._sectPr.find(qn('w:cols'))
    if node is None: node = OxmlElement('w:cols'); section._sectPr.append(node)
    for e in list(node): node.remove(e)
    node.set(qn('w:num'), str(count)); node.set(qn('w:space'), '482')
    node.set(qn('w:equalWidth'), '1')


def configure(doc, mode):
    paper = mode == 'paper'
    for s in doc.sections:
        s.page_width = Cm(21); s.page_height = Cm(29.7)
        s.left_margin = Cm(1.432 if paper else 3)
        s.right_margin = Cm(1.432 if paper else 2)
        s.top_margin = Cm(1.9 if paper else 2)
        s.bottom_margin = Cm(4.3 if paper else 2)
        s.header_distance = s.footer_distance = Cm(1)
    sizes = {'Normal': 10 if paper else 13, 'Title': 20 if paper else 18,
             'Heading 1': 10 if paper else 16, 'Heading 2': 10 if paper else 14,
             'Heading 3': 10 if paper else 13, 'Caption': 8.5 if paper else 11,
             'Reference': 8.5 if paper else 11, 'Index': 11, 'Small': 9 if paper else 10}
    for name, size in sizes.items():
        if name not in doc.styles: doc.styles.add_style(name, 1)
        style = doc.styles[name]; style.font.name = 'Times New Roman'; style.font.size = Pt(size)
        style.font.bold = name in ['Title', 'Heading 1', 'Heading 2', 'Heading 3']
        style.font.color.rgb = RGBColor(0, 0, 0)
        pf = style.paragraph_format
        pf.line_spacing = 1.0 if paper else (1.3 if name == 'Normal' else 1.1)
        pf.space_after = Pt(3 if paper else 6)
        pf.space_before = Pt((6 if paper else 8) if name.startswith('Heading') else 0)
        pf.widow_control = True
        pf.keep_with_next = name in ['Title', 'Heading 1', 'Heading 2', 'Heading 3']
        rf = style.element.get_or_add_rPr().get_or_add_rFonts()
        for k in list(rf.attrib):
            if 'theme' in k.lower(): del rf.attrib[k]
        for k in ['ascii', 'hAnsi', 'eastAsia', 'cs']: rf.set(qn('w:'+k), 'Times New Roman')
        for e in list(style.element.xpath('.//w:pBdr')): e.getparent().remove(e)
    if not paper:
        s = doc.sections[0]; s.different_first_page_header_footer = True
        p = s.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        font(p.add_run('Trang '), 10); field(p, 'PAGE')
    setting = OxmlElement('w:updateFields'); setting.set(qn('w:val'), 'true'); doc.settings.element.append(setting)


def add_table(doc, rows, mode, wide=False):
    n = len(rows[0]); paper = mode == 'paper'
    if any(len(r) != n for r in rows): raise ValueError('Inconsistent table columns')
    width = (18.136 if wide else 8.643) if paper else 16
    if rows[0][0] == 'Thành viên': weights = [1.2, 2.4, 1.7]
    elif rows[0][0] == 'Công trình': weights = [1.1, 1.2, 1.3, 2.5] if n == 4 else [1, 2, 2]
    elif rows[0][0] == 'Lượt': weights = [1.6, 1, 1, 1, 1, 1.3]
    elif n == 7 and rows[0][2] == 'Tham số': weights = [.7, .7, 1.3, 1, 1, 1, 1]
    elif n == 7: weights = [.8, .7, 1, 1, 1, .85, .85]
    elif rows[0][0] == 'Video': weights = [1.6, 1, 1]
    else: weights = [1]*n
    widths = [width*x/sum(weights) for x in weights]
    t = doc.add_table(rows=0, cols=n); t.autofit = False
    for col, w in zip(t.columns, widths): col.width = Cm(w)
    for i, values in enumerate(rows):
        row = t.add_row()
        for j, (cell, value, w) in enumerate(zip(row.cells, values, widths)):
            cell.width = Cm(w); cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]; pf = p.paragraph_format
            pf.first_line_indent = pf.left_indent = pf.right_indent = Cm(0)
            pf.line_spacing = 1; pf.space_before = pf.space_after = Pt(2)
            pf.keep_with_next = i == 0
            prose = len(value) > 40 or rows[0][0] in ['Công trình', 'Thành viên', 'Vị trí', 'Hạng mục']
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if prose else WD_ALIGN_PARAGRAPH.CENTER
            font(p.add_run(value), 8 if paper else 10.5, i == 0)
            cp = cell._tc.get_or_add_tcPr(); margins = OxmlElement('w:tcMar')
            for side in ['top', 'bottom', 'left', 'right']:
                e = OxmlElement('w:'+side); e.set(qn('w:w'), '45' if paper else '65'); e.set(qn('w:type'), 'dxa'); margins.append(e)
            cp.append(margins)
            sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), 'E7E6E6' if i == 0 else 'FFFFFF'); cp.append(sh)
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    borders = OxmlElement('w:tblBorders')
    for side in ['top', 'bottom', 'left', 'right', 'insideH', 'insideV']:
        e = OxmlElement('w:'+side)
        for k, v in {'val': 'single', 'sz': '4', 'color': 'D9D9D9'}.items(): e.set(qn('w:'+k), v)
        borders.append(e)
    t._tbl.tblPr.append(borders)
    return t


def add_equation(doc, text, paper=False):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    # Native Word math with Unicode symbols; no raw TeX or raster equations.
    def plain(parent, value):
        if not value: return
        r = OxmlElement('m:r'); t = OxmlElement('m:t'); t.text = value
        r.append(t); parent.append(r)
    def tokens(parent, value):
        power = re.search(r'(\([^)]*p_t\))\^γ', value)
        if power:
            tokens(parent, value[:power.start()])
            sup = OxmlElement('m:sSup'); base = OxmlElement('m:e'); exponent = OxmlElement('m:sup')
            tokens(base, power[1]); plain(exponent, 'γ'); sup.append(base); sup.append(exponent); parent.append(sup)
            tokens(parent, value[power.end():]); return
        offset = 0
        for m in re.finditer(r'([A-Za-zΑ-ω]+)_([A-Za-z]+)', value):
            plain(parent, value[offset:m.start()])
            sub = OxmlElement('m:sSub'); base = OxmlElement('m:e'); index = OxmlElement('m:sub')
            plain(base, m[1]); plain(index, m[2]); sub.append(base); sub.append(index); parent.append(sub)
            offset = m.end()
        plain(parent, value[offset:])
    math = OxmlElement('m:oMath'); tokens(math, text); p._p.append(math)


def build(mode, output=None, source=None):
    if mode not in OUTPUTS: raise ValueError(mode)
    source = Path(source or WORK/SOURCES[mode]).resolve()
    output = Path(output or WORK/OUTPUTS[mode]).resolve()
    if output in [x.resolve() for x in HISTORICAL] or output == source:
        raise ValueError('Historical deliverables and manuscript sources are read-only')
    old_hashes = {str(p.relative_to(ROOT)): sha(p) for p in HISTORICAL}
    text = source.read_text(encoding='utf8'); lines = text.splitlines()
    doc = Document(); configure(doc, mode); paper = mode == 'paper'
    tables = re.findall(r'^Bảng (\d+) (.+)$', text, re.M)
    figures = re.findall(r'^!\[Hình (\d+) ([^\]]+)\]\(([^)]+)\)$', text, re.M)
    in_cover = not paper; references = False; wide = False; image_inputs = []; i = 0
    while i < len(lines):
        line = lines[i].strip(); i += 1
        if not line: continue
        if line == '<!-- columns -->':
            columns(doc.add_section(WD_SECTION_START.CONTINUOUS), 2); continue
        if line in ['<!-- wide-start -->', '<!-- wide-end -->']:
            wide = line == '<!-- wide-start -->'
            columns(doc.add_section(WD_SECTION_START.CONTINUOUS), 1 if wide else 2)
            continue
        if line.startswith('<!--'): continue
        if line == '[[PAGE_BREAK]]': doc.add_page_break(); continue
        if line == '[[TOC]]':
            field(doc.add_paragraph(), r'TOC \o "1-2" \h \z \u'); continue
        if line in ['[[FIGURES]]', '[[TABLES]]']:
            for n, caption, *_ in (figures if line == '[[FIGURES]]' else tables):
                prefix = 'F' if line == '[[FIGURES]]' else 'T'
                p = doc.add_paragraph(style='Index')
                p.paragraph_format.tab_stops.add_tab_stop(Cm(16), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
                p.add_run(('Hình ' if prefix == 'F' else 'Bảng ')+n+' '+caption+'\t'); field(p, 'PAGEREF '+prefix+n+r' \h')
            continue
        if line.startswith('|'):
            raw = [line]
            while i < len(lines) and lines[i].strip().startswith('|'):
                raw.append(lines[i].strip()); i += 1
            rows = [[c.strip() for c in x.strip('|').split('|')] for x in raw if not re.fullmatch(r'\|[\s:|\-]+', x)]
            add_table(doc, rows, mode, wide); continue
        m = re.fullmatch(r'!\[Hình (\d+) ([^\]]+)\]\(([^)]+)\)', line)
        if m:
            n, caption, rel = m.groups(); path = (source.parent/rel).resolve()
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(path), width=Cm((17.8 if wide else 8.64) if paper else 16))
            p = doc.add_paragraph(style='Caption'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run('Hình '+n+' '+caption); bookmark(p, 'F'+n, 100+int(n))
            image_inputs.append({'path': str(path.relative_to(ROOT)), 'sha256': sha(path)})
            continue
        if line.startswith('::equation:: '): add_equation(doc, line[13:], paper); continue
        if re.match(r'^Bảng \d+ ', line):
            p = doc.add_paragraph(style='Caption'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = p.paragraph_format.keep_together = True
            p.add_run(line); n = re.match(r'Bảng (\d+)', line).group(1); bookmark(p, 'T'+n, 200+int(n)); continue
        if line.startswith('::author:: '):
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            font(p.add_run(line.split(":: ", 1)[1]), 11); continue
        if line.startswith(('::abstract:: ', '::keywords:: ', '::small:: ')):
            prefix, value = line.split(':: ', 1)
            p = doc.add_paragraph(style='Small' if 'small' in prefix or 'keywords' in prefix else 'Heading 2')
            p.add_run(value); continue
        heading = re.match(r'^(#{1,3}) (.+)$', line)
        if heading:
            level, value = len(heading[1]), heading[2]
            is_title = paper and level == 1 and not references and len(doc.paragraphs) == 0
            if not paper and level == 1:
                if value.startswith(('Phân công', 'Tóm tắt', 'Mục lục', 'Chương', 'Kết luận', 'Tài liệu', 'Phụ lục')):
                    doc.add_page_break(); in_cover = False
            if value == 'Tài liệu tham khảo': references = True
            style = 'Title' if is_title or in_cover else ('Heading 1' if level <= 2 and paper else 'Heading '+str(min(level, 3)))
            if not paper and not in_cover and level == 1 and not value.startswith(('Chương', 'Kết luận')): style = 'Title'
            if paper and level == 2:
                found = re.match(r'(\d+) (.+)', value)
                if found:
                    roman = ['I', 'II', 'III', 'IV', 'V', 'VI', 'VII'][int(found[1])-1]
                    value = roman+'. '+found[2].upper()
            elif paper and level == 3:
                value = re.sub(r'^\d+\.(\d+) ', lambda m: chr(64+int(m[1]))+'. ', value)
                style = 'Heading 2'
            p = doc.add_paragraph(value, style=style)
            if paper and level <= 2 or in_cover: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if not paper:
                bm = re.match(r'Chương (\d)', value)
                if bm: bookmark(p, 'Chuong'+bm[1], 900+int(bm[1]))
                elif value.startswith('Kết luận'): bookmark(p, 'KetLuan', 904)
                elif value.startswith('Tài liệu'): bookmark(p, 'TaiLieu', 905)
            continue
        p = doc.add_paragraph(style='Reference' if references else 'Normal')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if in_cover else (WD_ALIGN_PARAGRAPH.LEFT if references else WD_ALIGN_PARAGRAPH.JUSTIFY)
        p.add_run(line)
    # Geometry is inherited by all continuous sections. Explicitly audit it in tests.
    doc.core_properties.title = 'Đánh giá CDCN trên CASIA FASD đến lượt 7'
    doc.core_properties.author = 'Trịnh Minh Đạt; Nguyễn Hữu Minh Quân'
    doc.core_properties.subject = 'Báo cáo ba chương' if not paper else 'Bản thảo tiếng Việt theo bố cục tham chiếu IEEE conference'
    output.parent.mkdir(parents=True, exist_ok=True); doc.save(output)
    assert all(sha(ROOT/p) == h for p, h in old_hashes.items())
    provenance = {'source': str(source.relative_to(ROOT)) if source.is_relative_to(ROOT) else str(source),
                  'source_sha256': sha(source), 'output': output.name, 'output_sha256_before_field_refresh': sha(output),
                  'historical_read_only_sha256': old_hashes, 'images': image_inputs,
                  'evidence': {p: sha(ROOT/p) for p in EVIDENCE},
                  'scope': 'existing Rounds 1-7 evidence only; no training, scoring, webcam acquisition, slides or submission',
                  'paper_format': 'IEEE conference reference layout, A4 two columns 10pt; Vietnamese adaptation, no target venue' if paper else None}
    output.with_suffix('.provenance.json').write_text(json.dumps(provenance, ensure_ascii=False, indent=2)+'\n', encoding='utf8')
    return output


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--kind', choices=['report', 'paper', 'both'], default='both')
    args = parser.parse_args()
    for mode in (['report', 'paper'] if args.kind == 'both' else [args.kind]): print(build(mode))
